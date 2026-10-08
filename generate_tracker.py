#!/usr/bin/env python3
"""
Full Stack + GenAI Roadmap — Progress Tracker Excel Generator
Run: python3 generate_tracker.py
"""

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import Rule, DataBarRule
from openpyxl.styles.differential import DifferentialStyle

OUTPUT = "/Users/durgesh.rathod/Documents/personal/roadmaps/roadmap-tracker.xlsx"

# ─────────────────────────────────────────────────────────────────────────────
# STYLE HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")

def font(bold=False, size=11, color="000000", italic=False, name="Calibri"):
    return Font(bold=bold, size=size, color=color, italic=italic, name=name)

def align(h="center", v="center", wrap=True):
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def border(color="CCCCCC", style="thin"):
    s = Side(style=style, color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def thick_border(color="888888"):
    s = Side(style="medium", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def cw(ws, col_idx, width):
    ws.column_dimensions[get_column_letter(col_idx)].width = width

def rh(ws, row_idx, height):
    ws.row_dimensions[row_idx].height = height

def write(ws, row, col, value, bold=False, size=11, fc="000000", bg=None,
          ha="left", va="center", wrap=True, italic=False, bc="CCCCCC",
          num_fmt=None):
    c = ws.cell(row=row, column=col)
    c.value = value
    c.font = font(bold=bold, size=size, color=fc, italic=italic)
    c.alignment = align(h=ha, v=va, wrap=wrap)
    c.border = border(color=bc)
    if bg:
        c.fill = fill(bg)
    if num_fmt:
        c.number_format = num_fmt
    return c

def merge_write(ws, r1, c1, r2, c2, value, bold=False, size=11, fc="000000",
                bg=None, ha="center", va="center", bc="888888"):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    c = ws.cell(row=r1, column=c1)
    c.value = value
    c.font = font(bold=bold, size=size, color=fc)
    c.alignment = align(h=ha, v=va)
    c.border = thick_border(bc)
    if bg:
        c.fill = fill(bg)
    return c

def add_dropdown(ws, formula, sqref):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showDropDown=False)
    dv.sqref = sqref
    ws.add_data_validation(dv)

def add_cf_status(ws, col_letter, data_range):
    """Conditional formatting: color rows based on status column value."""
    dxf_done    = DifferentialStyle(fill=PatternFill(bgColor="C8E6C9", fill_type="solid"), font=Font(color="1B5E20"))
    dxf_inprog  = DifferentialStyle(fill=PatternFill(bgColor="FFF9C4", fill_type="solid"), font=Font(color="F57F17"))
    dxf_blocked = DifferentialStyle(fill=PatternFill(bgColor="FFCDD2", fill_type="solid"), font=Font(color="B71C1C"))
    dxf_skipped = DifferentialStyle(fill=PatternFill(bgColor="E0E0E0", fill_type="solid"), font=Font(color="616161"))

    ws.conditional_formatting.add(data_range,
        Rule(type="expression", formula=[f'${col_letter}1="Done"'], dxf=dxf_done, stopIfTrue=True))
    ws.conditional_formatting.add(data_range,
        Rule(type="expression", formula=[f'${col_letter}1="In Progress"'], dxf=dxf_inprog, stopIfTrue=True))
    ws.conditional_formatting.add(data_range,
        Rule(type="expression", formula=[f'${col_letter}1="Blocked"'], dxf=dxf_blocked, stopIfTrue=True))
    ws.conditional_formatting.add(data_range,
        Rule(type="expression", formula=[f'${col_letter}1="Skipped"'], dxf=dxf_skipped, stopIfTrue=True))

# ─────────────────────────────────────────────────────────────────────────────
# DATA
# ─────────────────────────────────────────────────────────────────────────────

PHASES = [
    {"name": "Phase 1 — Foundation",        "color": "1976D2", "light": "DDEEFF",
     "weeks": "1–3",   "objective": "Learn JS basics + build first REST API"},
    {"name": "Phase 2 — Backend",           "color": "00897B", "light": "D4F0EC",
     "weeks": "4–6",   "objective": "Build real backend with DB + Auth"},
    {"name": "Phase 3 — Full Stack",        "color": "43A047", "light": "DCF0DC",
     "weeks": "7–9",   "objective": "Build complete UI + backend application"},
    {"name": "Phase 4 — GenAI",             "color": "8E24AA", "light": "F0E0F5",
     "weeks": "10–12", "objective": "Add LLM, embeddings, RAG pipeline"},
    {"name": "Phase 5 — Interview Project", "color": "3949AB", "light": "E0E4F5",
     "weeks": "13–16", "objective": "Build Document QnA Chatbot from scratch"},
    {"name": "Phase 6 — Deployment + Jobs", "color": "FB8C00", "light": "FFF0D4",
     "weeks": "—",     "objective": "Deploy publicly + Resume + Apply"},
]

# (phase, #, weeks, category, task, description)
TASKS = [
    # ── Phase 1 ──────────────────────────────────────────────────────────────
    ("Phase 1 — Foundation","1","1–3","Setup","Install Node.js & VS Code",
     "Install Node.js LTS, VS Code, ESLint + Prettier extensions, configure workspace"),
    ("Phase 1 — Foundation","2","1–3","JS Basics","Variables & Data Types",
     "var/let/const, string/number/boolean/null/undefined, typeof, type coercion"),
    ("Phase 1 — Foundation","3","1–3","JS Basics","Control Flow",
     "if/else, switch, ternary, for/while/do-while, break/continue"),
    ("Phase 1 — Foundation","4","1–3","JS Basics","Arrays & Methods",
     "map, filter, reduce, find, forEach, some, every, flat, spread operator"),
    ("Phase 1 — Foundation","5","1–3","JS Basics","Objects & Prototypes",
     "Literals, destructuring, optional chaining (?.), nullish coalescing (??), spread"),
    ("Phase 1 — Foundation","6","1–3","JS Basics","Functions & Closures",
     "Arrow functions, default params, rest params, IIFE, closures, HOFs"),
    ("Phase 1 — Foundation","7","1–3","Async JS","Callbacks & Promises",
     "Callback hell, Promise chain (.then/.catch), Promise.all, Promise.race"),
    ("Phase 1 — Foundation","8","1–3","Async JS","async / await",
     "async functions, await keyword, try/catch error handling, parallel await"),
    ("Phase 1 — Foundation","9","1–3","ES6+","Modules & Modern Syntax",
     "import/export, template literals, destructuring assignment, optional chaining"),
    ("Phase 1 — Foundation","10","1–3","Express.js","Project Setup",
     "npm init, install express cors body-parser, nodemon, folder structure (routes/controllers)"),
    ("Phase 1 — Foundation","11","1–3","Express.js","GET Endpoints",
     "app.get(), route params (:id), query strings (?key=val), send JSON responses"),
    ("Phase 1 — Foundation","12","1–3","Express.js","POST Endpoints",
     "app.post(), express.json() middleware, body parsing, create resource flow"),
    ("Phase 1 — Foundation","13","1–3","Express.js","Middleware & Errors",
     "Custom middleware chain, cors(), morgan, 404 handler, global error middleware"),
    ("Phase 1 — Foundation","14","1–3","Testing","Postman Collections",
     "Create workspace, collection, test GET/POST/PUT/DELETE, environment vars, export"),
    ("Phase 1 — Foundation","15","1–3","Testing","REST Principles",
     "HTTP methods, status codes (200/201/400/401/404/500), idempotency, RESTful naming"),
    # ── Phase 2 ──────────────────────────────────────────────────────────────
    ("Phase 2 — Backend","1","4–6","PostgreSQL","Install & Configure",
     "Install PostgreSQL, create DB user + database, connect via psql or pgAdmin"),
    ("Phase 2 — Backend","2","4–6","PostgreSQL","SQL Fundamentals",
     "CREATE TABLE, INSERT, SELECT, WHERE, JOIN (INNER/LEFT), ORDER BY, LIMIT"),
    ("Phase 2 — Backend","3","4–6","PostgreSQL","DB Schema Design",
     "Design users & tasks tables, primary keys, foreign keys, constraints, indexes"),
    ("Phase 2 — Backend","4","4–6","Prisma ORM","Setup & Configure",
     "npm install prisma @prisma/client, prisma init, configure DATABASE_URL in .env"),
    ("Phase 2 — Backend","5","4–6","Prisma ORM","Schema Definition",
     "Define User and Task models, @relation, @unique, @default, DateTime fields"),
    ("Phase 2 — Backend","6","4–6","Prisma ORM","Migrations",
     "prisma migrate dev --name init, inspect generated SQL, prisma studio"),
    ("Phase 2 — Backend","7","4–6","Prisma ORM","Prisma Client CRUD",
     "findMany, findUnique, create, update, delete, upsert with Prisma Client"),
    ("Phase 2 — Backend","8","4–6","Auth","Password Hashing (bcrypt)",
     "npm install bcryptjs, hash on register (saltRounds=10), compare on login"),
    ("Phase 2 — Backend","9","4–6","Auth","JWT Token Generation",
     "npm install jsonwebtoken, sign with secret+expiry, verify, decode payload"),
    ("Phase 2 — Backend","10","4–6","Auth","Auth Middleware",
     "Extract Bearer token from header, verify JWT, attach req.user, protect routes"),
    ("Phase 2 — Backend","11","4–6","Auth","Register Endpoint",
     "POST /auth/register — validate input, check duplicate email, hash, save to DB"),
    ("Phase 2 — Backend","12","4–6","Auth","Login Endpoint",
     "POST /auth/login — find user, compare hash, sign + return JWT + user data"),
    ("Phase 2 — Backend","13","4–6","CRUD API","Task Endpoints",
     "GET/POST/PUT/DELETE /tasks — all protected, user-scoped, full CRUD with DB"),
    ("Phase 2 — Backend","14","4–6","CRUD API","Input Validation",
     "Validate required fields, types, lengths; return 400 with descriptive errors"),
    ("Phase 2 — Backend","15","4–6","CRUD API","Integration Test",
     "Postman: register → login → create/read/update/delete tasks → verify auth guards"),
    # ── Phase 3 ──────────────────────────────────────────────────────────────
    ("Phase 3 — Full Stack","1","7–9","React","Project Setup (Vite)",
     "npm create vite@latest, choose React template, clean boilerplate, folder structure"),
    ("Phase 3 — Full Stack","2","7–9","React","Components & Props",
     "Functional components, prop types, children, composition, reusable UI components"),
    ("Phase 3 — Full Stack","3","7–9","React","useState Hook",
     "Form state, counter, toggle, controlled inputs, derived state patterns"),
    ("Phase 3 — Full Stack","4","7–9","React","useEffect Hook",
     "Fetch on mount, cleanup function, dependency array, avoid infinite re-renders"),
    ("Phase 3 — Full Stack","5","7–9","React","Custom Hooks",
     "useAuth (token/user state), useTasks (fetch + CRUD), useForm (generic handler)"),
    ("Phase 3 — Full Stack","6","7–9","Tailwind CSS","Setup & Config",
     "npm install tailwindcss, tailwind.config.js, JIT mode, purge paths, base styles"),
    ("Phase 3 — Full Stack","7","7–9","Tailwind CSS","Layout & Responsive",
     "Flexbox, grid, responsive breakpoints (sm/md/lg/xl), sidebar + main layout"),
    ("Phase 3 — Full Stack","8","7–9","Tailwind CSS","Theming & Reuse",
     "Custom colors in config, dark mode, @apply for repeated component styles"),
    ("Phase 3 — Full Stack","9","7–9","Axios","Instance & Interceptors",
     "Create Axios instance with baseURL, request interceptor to attach JWT header"),
    ("Phase 3 — Full Stack","10","7–9","Axios","API Service Layer",
     "auth.js and tasks.js service files — centralise all API calls, handle errors"),
    ("Phase 3 — Full Stack","11","7–9","Auth UI","Login & Register Forms",
     "Controlled forms, client-side validation, show API errors inline, loading button"),
    ("Phase 3 — Full Stack","12","7–9","Auth UI","Token Storage",
     "Store JWT in localStorage, parse on app load, clear on logout, expiry handling"),
    ("Phase 3 — Full Stack","13","7–9","Auth UI","Protected Routes",
     "React Router v6, <RequireAuth> wrapper, redirect to /login if no valid token"),
    ("Phase 3 — Full Stack","14","7–9","Task UI","Task List",
     "Fetch & render tasks, loading skeleton, empty state illustration, task count"),
    ("Phase 3 — Full Stack","15","7–9","Task UI","Task CRUD UI",
     "Add task form, inline edit mode, delete confirm dialog, complete toggle checkbox"),
    ("Phase 3 — Full Stack","16","7–9","UX","Loading & Error States",
     "Spinner/skeleton during fetch, toast notifications (react-hot-toast), retry button"),
    ("Phase 3 — Full Stack","17","7–9","UX","End-to-End Integration Test",
     "Register → Login → CRUD tasks → Logout → all UI states verified manually"),
    # ── Phase 4 ──────────────────────────────────────────────────────────────
    ("Phase 4 — GenAI","1","10–12","Ollama","Install & Start Service",
     "brew install ollama, ollama serve, verify running on localhost:11434"),
    ("Phase 4 — GenAI","2","10–12","Ollama","Pull LLM Models",
     "ollama pull llama3, ollama pull mistral, compare output quality"),
    ("Phase 4 — GenAI","3","10–12","Ollama","Test LLM via Terminal",
     "ollama run llama3, chat, try different prompts, understand context window"),
    ("Phase 4 — GenAI","4","10–12","Ollama","Ollama REST API",
     "POST /api/generate, stream:false, stream:true — call from Node.js with fetch"),
    ("Phase 4 — GenAI","5","10–12","Embeddings","Concept & Theory",
     "What are embeddings, vector space, cosine similarity, why they matter for RAG"),
    ("Phase 4 — GenAI","6","10–12","Embeddings","Pull Embedding Model",
     "ollama pull nomic-embed-text, test POST /api/embeddings, inspect vector dimensions"),
    ("Phase 4 — GenAI","7","10–12","Embeddings","Similarity Search Concept",
     "Compute cosine similarity manually in JS, understand k-nearest-neighbor"),
    ("Phase 4 — GenAI","8","10–12","LangChain JS","Setup",
     "npm install langchain @langchain/community @langchain/ollama"),
    ("Phase 4 — GenAI","9","10–12","LangChain JS","Simple Chain",
     "PromptTemplate → OllamaLLM → StringOutputParser — test with simple questions"),
    ("Phase 4 — GenAI","10","10–12","FAISS","Setup & Basic Ops",
     "npm install faiss-node (or via @langchain/community), create store, basic test"),
    ("Phase 4 — GenAI","11","10–12","FAISS","Vector Store CRUD",
     "addDocuments, similaritySearch(query, k), saveLocal, loadFromLocal"),
    ("Phase 4 — GenAI","12","10–12","RAG","Build RAG Chain",
     "Query → FAISS similarity search → format context → Ollama LLM → answer"),
    ("Phase 4 — GenAI","13","10–12","RAG","End-to-End Test",
     "Feed 5+ documents, ask 10 questions, verify retrieved context is relevant"),
    ("Phase 4 — GenAI","14","10–12","AI Features","NL Task Input",
     "Parse natural language (\"remind me Friday\") via LLM → structured task → DB"),
    ("Phase 4 — GenAI","15","10–12","AI Features","AI Task Suggestions",
     "Based on existing tasks, ask LLM for next recommended actions with context"),
    ("Phase 4 — GenAI","16","10–12","AI Features","Full Integration Test",
     "Test NL input + AI suggestions in TODO app (React + Express + Ollama + FAISS)"),
    # ── Phase 5 ──────────────────────────────────────────────────────────────
    ("Phase 5 — Interview Project","1","13–16","Planning","Architecture Design",
     "Draw full system diagram, define all API endpoints, data flow, tech decisions"),
    ("Phase 5 — Interview Project","2","13–16","Backend","Express.js Setup",
     "New project, folder structure (routes/controllers/services), env vars, nodemon"),
    ("Phase 5 — Interview Project","3","13–16","Backend","File Upload (Multer)",
     "npm install multer, diskStorage config, MIME type + size validation, /uploads dir"),
    ("Phase 5 — Interview Project","4","13–16","Backend","PDF Parser",
     "LangChain PDFLoader or pdf-parse — extract raw text from uploaded PDF"),
    ("Phase 5 — Interview Project","5","13–16","Backend","TXT Parser",
     "fs.readFile for plain text, normalize whitespace, handle encoding (UTF-8)"),
    ("Phase 5 — Interview Project","6","13–16","LangChain","Text Chunking",
     "RecursiveCharacterTextSplitter, chunkSize:500, chunkOverlap:50, inspect chunks"),
    ("Phase 5 — Interview Project","7","13–16","LangChain","Embed Chunks",
     "OllamaEmbeddings with nomic-embed-text, embed all chunks, inspect vector shape"),
    ("Phase 5 — Interview Project","8","13–16","FAISS","Build Vector Store",
     "FaissStore.fromDocuments(), save per documentId, load on query request"),
    ("Phase 5 — Interview Project","9","13–16","Backend","Retrieval Endpoint",
     "POST /ask — load FAISS for doc, similaritySearch top 5, return chunks + metadata"),
    ("Phase 5 — Interview Project","10","13–16","LangChain","RAG Pipeline",
     "RetrievalQAChain or custom: retrieve → format context → Ollama → response"),
    ("Phase 5 — Interview Project","11","13–16","Backend","Streaming Response",
     "Stream Ollama response via Server-Sent Events (res.write) or chunked transfer"),
    ("Phase 5 — Interview Project","12","13–16","Frontend","React + Tailwind Setup",
     "Vite + React, install Tailwind + Axios, folder structure, Axios base instance"),
    ("Phase 5 — Interview Project","13","13–16","Frontend","Upload Panel",
     "Drag-and-drop or click-to-upload, file type badge, upload progress indicator"),
    ("Phase 5 — Interview Project","14","13–16","Frontend","Chat UI",
     "Message thread (user/bot bubbles), auto-scroll to bottom, timestamps, avatars"),
    ("Phase 5 — Interview Project","15","13–16","Frontend","Streaming in UI",
     "Handle chunked tokens as they arrive, typewriter effect, stop generation button"),
    ("Phase 5 — Interview Project","16","13–16","Frontend","Loading & Error States",
     "Skeleton loaders, error toast, retry button, empty state with upload prompt"),
    ("Phase 5 — Interview Project","17","13–16","Quality","Full Local Test",
     "Upload 3+ docs, ask 10+ questions per doc, verify answer accuracy and relevance"),
    ("Phase 5 — Interview Project","18","13–16","Quality","Write README",
     "Architecture diagram, setup instructions, screenshots, demo GIF, tech stack badges"),
    # ── Phase 6 ──────────────────────────────────────────────────────────────
    ("Phase 6 — Deployment + Jobs","1","—","Deployment","Vercel Account Setup",
     "Sign up Vercel, install CLI (npm i -g vercel), link GitHub account"),
    ("Phase 6 — Deployment + Jobs","2","—","Deployment","Frontend Env Vars",
     "Set VITE_API_URL to Render backend URL in Vercel Environment Variables settings"),
    ("Phase 6 — Deployment + Jobs","3","—","Deployment","Deploy Frontend to Vercel",
     "vercel --prod, verify build output, test on *.vercel.app live URL"),
    ("Phase 6 — Deployment + Jobs","4","—","Deployment","Render Account Setup",
     "Sign up Render, New Web Service, connect GitHub repo, set runtime: Node"),
    ("Phase 6 — Deployment + Jobs","5","—","Deployment","Backend Env Vars",
     "Set DATABASE_URL, JWT_SECRET, CORS_ORIGIN, NODE_ENV=production on Render"),
    ("Phase 6 — Deployment + Jobs","6","—","Deployment","Deploy Backend to Render",
     "Set build cmd: npm install, start cmd: node server.js, enable auto-deploy"),
    ("Phase 6 — Deployment + Jobs","7","—","Deployment","Cloud Database (Neon)",
     "Create Neon or Supabase project, copy connection string, set on Render env"),
    ("Phase 6 — Deployment + Jobs","8","—","Deployment","Migrate Schema to Cloud",
     "DATABASE_URL=<neon_url> npx prisma migrate deploy — verify tables in cloud"),
    ("Phase 6 — Deployment + Jobs","9","—","Deployment","Smoke Test Live App",
     "Login, CRUD tasks, upload document, ask question — all on production URL"),
    ("Phase 6 — Deployment + Jobs","10","—","Portfolio","Public URL + Screenshots",
     "Copy live URL, take screenshots of key flows, record 60-sec demo GIF"),
    ("Phase 6 — Deployment + Jobs","11","—","Portfolio","Update README",
     "Add live demo badge, screenshots, architecture diagram, local setup guide"),
    ("Phase 6 — Deployment + Jobs","12","—","Portfolio","GitHub Cleanup",
     "Pin 2 repos, add descriptions, clean commits, remove test files + console.logs"),
    ("Phase 6 — Deployment + Jobs","13","—","Job Prep","Resume Bullets",
     "2-3 bullets per project: action verb + tech stack + measurable outcome/metric"),
    ("Phase 6 — Deployment + Jobs","14","—","Job Prep","LinkedIn Profile Update",
     "Update headline + about section, add projects, post publicly about project"),
    ("Phase 6 — Deployment + Jobs","15","—","Job Prep","Apply Round 1",
     "Apply to 10 jobs on LinkedIn + Naukri with tailored resume + cover note"),
    ("Phase 6 — Deployment + Jobs","16","—","Job Prep","Apply Round 2",
     "10 more applications, ask connections for referrals, follow up after 1 week"),
    ("Phase 6 — Deployment + Jobs","17","—","Interview Prep","Mock: TODO App System",
     "Explain architecture, DB schema design, auth flow, Prisma decisions, trade-offs"),
    ("Phase 6 — Deployment + Jobs","18","—","Interview Prep","Mock: Doc QnA AI Flow",
     "Explain RAG pipeline step-by-step, embedding strategy, why FAISS, LangChain role"),
    ("Phase 6 — Deployment + Jobs","19","—","Interview Prep","DSA Prep",
     "Arrays, strings, hashmaps — solve 20 easy LeetCode problems, time yourself"),
    ("Phase 6 — Deployment + Jobs","20","—","Interview Prep","HR Questions",
     "Tell me about yourself, strengths, why this role, salary expectations, notice period"),
]

# ─────────────────────────────────────────────────────────────────────────────
# SHEET 1: DASHBOARD
# ─────────────────────────────────────────────────────────────────────────────

def build_dashboard(wb, tt_sheet_name="Task Tracker"):
    ws = wb.create_sheet("Dashboard")
    ws.sheet_properties.tabColor = "1976D2"
    ws.sheet_view.showGridLines = False

    # Column widths
    widths = [2, 30, 10, 42, 32, 12, 10, 13, 11, 14, 24, 14, 2]
    for i, w in enumerate(widths, 1):
        cw(ws, i, w)

    # ── Title ──────────────────────────────────────────────────────────────
    rh(ws, 1, 10)
    rh(ws, 2, 52)
    merge_write(ws, 2, 2, 2, 12, "Full Stack + GenAI  —  Roadmap Progress Tracker",
                bold=True, size=22, fc="FFFFFF", bg="1A237E", ha="center")
    rh(ws, 3, 28)
    merge_write(ws, 3, 2, 3, 12, "Track your journey from zero to Full Stack + GenAI engineer",
                bold=False, size=12, fc="E8EAF6", bg="283593", ha="center", bc="283593")

    # ── Summary row ────────────────────────────────────────────────────────
    rh(ws, 4, 10)
    rh(ws, 5, 34)
    # Overall % (formula-based)
    write(ws, 5, 2, "Overall Progress", bold=True, size=12, fc="1A237E", ha="left", bc="E0E0E0")
    c = ws.cell(row=5, column=3)
    c.value = f"=IFERROR(COUNTIF('{tt_sheet_name}'!H:H,\"Done\")/COUNTA('{tt_sheet_name}'!H:H),0)"
    c.number_format = "0%"
    c.font = font(bold=True, size=14, color="1976D2")
    c.alignment = align("center", "center")
    c.border = border("CCCCCC")

    write(ws, 5, 4, "=TODAY()", bold=False, size=11, fc="616161", ha="left",
          bc="E0E0E0", num_fmt="DD MMM YYYY")
    ws.cell(row=5, column=4).number_format = "DD MMM YYYY"
    rh(ws, 6, 8)

    # ── Phase table headers ─────────────────────────────────────────────────
    rh(ws, 7, 30)
    header_bg = "37474F"
    headers = ["Phase", "Weeks", "Objective", "Project", "Total", "Done", "In Prog", "Blocked", "% Done", "Progress Bar", "Current Status"]
    for col_i, h in enumerate(headers, 2):
        write(ws, 7, col_i, h, bold=True, size=11, fc="FFFFFF", bg=header_bg,
              ha="center", bc="546E7A")

    # ── Phase rows ─────────────────────────────────────────────────────────
    projects = [
        "—",
        "TODO App (Backend)",
        "TODO App (Full Stack)",
        "TODO App + AI",
        "Document QnA Chatbot",
        "Live on Vercel + Render",
    ]

    for idx, (ph, proj) in enumerate(zip(PHASES, projects), 1):
        r = 7 + idx
        rh(ws, r, 38)
        nm = ph["name"]
        bg = ph["light"]
        hc = ph["color"]

        write(ws, r, 2, nm, bold=True, size=11, fc="FFFFFF", bg=hc, ha="left", bc=hc)
        write(ws, r, 3, ph["weeks"], bold=False, size=11, fc="000000", bg=bg, ha="center", bc="CCCCCC")
        write(ws, r, 4, ph["objective"], bold=False, size=10, fc="333333", bg=bg, ha="left", bc="CCCCCC")
        write(ws, r, 5, proj, bold=False, size=10, fc="333333", bg=bg, ha="left", bc="CCCCCC")

        # Formulas counting from Task Tracker
        total_f  = f'=COUNTIF(\'{tt_sheet_name}\'!B:B,"{nm}")'
        done_f   = f'=COUNTIFS(\'{tt_sheet_name}\'!B:B,"{nm}",\'{tt_sheet_name}\'!H:H,"Done")'
        inprog_f = f'=COUNTIFS(\'{tt_sheet_name}\'!B:B,"{nm}",\'{tt_sheet_name}\'!H:H,"In Progress")'
        blkd_f   = f'=COUNTIFS(\'{tt_sheet_name}\'!B:B,"{nm}",\'{tt_sheet_name}\'!H:H,"Blocked")'
        pct_f    = f"=IFERROR(G{r}/F{r},0)"
        bar_f    = f'=IFERROR(REPT("█",ROUND(J{r}*20,0))&REPT("░",20-ROUND(J{r}*20,0))&"  "&TEXT(J{r},"0%"),"—")'
        status_f = f'=IF(J{r}=1,"✅ Complete",IF(J{r}=0,"⬜ Not Started",IF(K{r}>0,"🔴 Blocked","🔵 In Progress")))'

        for col_i, fml in [(6, total_f), (7, done_f), (8, inprog_f), (9, blkd_f)]:
            c = ws.cell(row=r, column=col_i)
            c.value = fml
            c.fill = fill(bg)
            c.font = font(bold=True, size=11, color="000000")
            c.alignment = align("center", "center")
            c.border = border("CCCCCC")

        pct_cell = ws.cell(row=r, column=10)
        pct_cell.value = pct_f
        pct_cell.number_format = "0%"
        pct_cell.fill = fill(bg)
        pct_cell.font = font(bold=True, size=11, color=hc)
        pct_cell.alignment = align("center", "center")
        pct_cell.border = border("CCCCCC")

        bar_cell = ws.cell(row=r, column=11)
        bar_cell.value = bar_f
        bar_cell.fill = fill("F5F5F5")
        bar_cell.font = Font(name="Consolas", size=10, color="1976D2")
        bar_cell.alignment = align("left", "center", wrap=False)
        bar_cell.border = border("CCCCCC")

        status_cell = ws.cell(row=r, column=12)
        status_cell.value = status_f
        status_cell.fill = fill(bg)
        status_cell.font = font(size=11, color="333333")
        status_cell.alignment = align("center", "center")
        status_cell.border = border("CCCCCC")

    # ── Totals row ─────────────────────────────────────────────────────────
    tr = 14
    rh(ws, tr, 32)
    merge_write(ws, tr, 2, tr, 5, "TOTAL", bold=True, size=12, fc="FFFFFF", bg="37474F")
    for col_i, rng in [(6, "F8:F13"), (7, "G8:G13"), (8, "H8:H13"), (9, "I8:I13")]:
        c = ws.cell(row=tr, column=col_i)
        c.value = f"=SUM({rng})"
        c.fill = fill("455A64")
        c.font = font(bold=True, size=12, color="FFFFFF")
        c.alignment = align("center", "center")
        c.border = border("546E7A")

    overall_pct = ws.cell(row=tr, column=10)
    overall_pct.value = f"=IFERROR(G{tr}/F{tr},0)"
    overall_pct.number_format = "0%"
    overall_pct.fill = fill("455A64")
    overall_pct.font = font(bold=True, size=12, color="76FF03")
    overall_pct.alignment = align("center", "center")
    overall_pct.border = border("546E7A")

    overall_bar = ws.cell(row=tr, column=11)
    overall_bar.value = f'=REPT("█",ROUND(J{tr}*20,0))&REPT("░",20-ROUND(J{tr}*20,0))&"  "&TEXT(J{tr},"0%")'
    overall_bar.fill = fill("37474F")
    overall_bar.font = Font(name="Consolas", size=11, color="76FF03", bold=True)
    overall_bar.alignment = align("left", "center", wrap=False)
    overall_bar.border = border("546E7A")

    # ── Legend ─────────────────────────────────────────────────────────────
    rh(ws, 15, 10)
    rh(ws, 16, 24)
    merge_write(ws, 16, 2, 16, 12, "STATUS LEGEND", bold=True, size=11, fc="FFFFFF", bg="546E7A")
    legend_items = [
        ("Not Started", "F5F5F5", "616161"),
        ("In Progress",  "FFF9C4", "F57F17"),
        ("Done",         "C8E6C9", "1B5E20"),
        ("Blocked",      "FFCDD2", "B71C1C"),
        ("Skipped",      "E0E0E0", "424242"),
    ]
    rh(ws, 17, 26)
    for col_i, (label, bg_c, fc_c) in enumerate(legend_items, 2):
        write(ws, 17, col_i, label, bold=True, size=11, fc=fc_c, bg=bg_c, ha="center", bc="AAAAAA")

    return ws

# ─────────────────────────────────────────────────────────────────────────────
# SHEET 2: TASK TRACKER
# ─────────────────────────────────────────────────────────────────────────────

def build_task_tracker(wb):
    ws = wb.create_sheet("Task Tracker")
    ws.sheet_properties.tabColor = "43A047"
    ws.sheet_view.showGridLines = False

    # Column widths: A #, B Phase, C Weeks, D Category, E Task, F Description,
    #                G Priority, H Status, I Start, J Target, K Done, L Hours, M Notes
    widths = [5, 28, 7, 18, 28, 45, 10, 14, 12, 12, 12, 9, 35]
    for i, w in enumerate(widths, 1):
        cw(ws, i, w)

    # ── Header ─────────────────────────────────────────────────────────────
    rh(ws, 1, 42)
    merge_write(ws, 1, 1, 1, 13, "Task Tracker — Full Stack + GenAI Roadmap",
                bold=True, size=18, fc="FFFFFF", bg="1A237E", ha="center")
    rh(ws, 2, 8)

    # Column headers row
    rh(ws, 3, 30)
    col_headers = ["#", "Phase", "Wks", "Category", "Task", "Description",
                   "Priority", "Status", "Start Date", "Target Date", "Done Date", "Hrs", "Notes / Learnings"]
    header_bg = "37474F"
    for col_i, h in enumerate(col_headers, 1):
        write(ws, 3, col_i, h, bold=True, size=11, fc="FFFFFF", bg=header_bg, ha="center", bc="546E7A")

    ws.freeze_panes = "A4"

    # ── Data rows ──────────────────────────────────────────────────────────
    phase_color_map = {ph["name"]: ph for ph in PHASES}
    row = 4
    task_row_start = 4

    for t in TASKS:
        phase_name, num, weeks, category, task_name, desc = t
        ph = phase_color_map[phase_name]
        hc = ph["color"]
        bg = ph["light"]
        rh(ws, row, 34)

        write(ws, row, 1,  num,         bold=True,  size=10, fc="FFFFFF", bg=hc, ha="center", bc=hc)
        write(ws, row, 2,  phase_name,  bold=True,  size=10, fc="FFFFFF", bg=hc, ha="left",   bc=hc)
        write(ws, row, 3,  weeks,       bold=False, size=10, fc="444444", bg=bg, ha="center", bc="CCCCCC")
        write(ws, row, 4,  category,    bold=False, size=10, fc="444444", bg=bg, ha="left",   bc="CCCCCC")
        write(ws, row, 5,  task_name,   bold=True,  size=10, fc="222222", bg=bg, ha="left",   bc="CCCCCC")
        write(ws, row, 6,  desc,        bold=False, size=9,  fc="444444", bg="FAFAFA", ha="left", bc="CCCCCC", wrap=True)
        write(ws, row, 7,  "Medium",    bold=False, size=10, fc="555555", bg="FFFFFF", ha="center", bc="CCCCCC")
        write(ws, row, 8,  "Not Started",bold=False,size=10, fc="777777", bg="F5F5F5", ha="center", bc="CCCCCC")
        write(ws, row, 9,  None,        bold=False, size=10, fc="555555", bg="FFFFFF", ha="center", bc="CCCCCC", num_fmt="DD-MMM-YY")
        write(ws, row, 10, None,        bold=False, size=10, fc="555555", bg="FFFFFF", ha="center", bc="CCCCCC", num_fmt="DD-MMM-YY")
        write(ws, row, 11, None,        bold=False, size=10, fc="555555", bg="FFFFFF", ha="center", bc="CCCCCC", num_fmt="DD-MMM-YY")
        write(ws, row, 12, None,        bold=False, size=10, fc="555555", bg="FFFFFF", ha="center", bc="CCCCCC")
        write(ws, row, 13, "",          bold=False, size=9,  fc="555555", bg="FFFDE7", ha="left",   bc="CCCCCC")
        row += 1

    task_row_end = row - 1

    # Dropdowns
    add_dropdown(ws, '"Not Started,In Progress,Done,Blocked,Skipped"', f"H4:H{task_row_end}")
    add_dropdown(ws, '"High,Medium,Low"',                               f"G4:G{task_row_end}")

    # Conditional formatting on entire row based on Status column (H)
    data_rng = f"A4:M{task_row_end}"
    dxf_done    = DifferentialStyle(fill=PatternFill(bgColor="C8E6C9", fill_type="solid"), font=Font(color="1B5E20", name="Calibri"))
    dxf_inprog  = DifferentialStyle(fill=PatternFill(bgColor="FFF9C4", fill_type="solid"), font=Font(color="F57F17", name="Calibri"))
    dxf_blocked = DifferentialStyle(fill=PatternFill(bgColor="FFCDD2", fill_type="solid"), font=Font(color="B71C1C", name="Calibri"))
    dxf_skipped = DifferentialStyle(fill=PatternFill(bgColor="EEEEEE", fill_type="solid"), font=Font(color="9E9E9E", name="Calibri", italic=True))

    ws.conditional_formatting.add(data_rng, Rule(type="expression", formula=['$H4="Done"'], dxf=dxf_done, stopIfTrue=True))
    ws.conditional_formatting.add(data_rng, Rule(type="expression", formula=['$H4="In Progress"'], dxf=dxf_inprog, stopIfTrue=True))
    ws.conditional_formatting.add(data_rng, Rule(type="expression", formula=['$H4="Blocked"'], dxf=dxf_blocked, stopIfTrue=True))
    ws.conditional_formatting.add(data_rng, Rule(type="expression", formula=['$H4="Skipped"'], dxf=dxf_skipped, stopIfTrue=True))

    return ws

# ─────────────────────────────────────────────────────────────────────────────
# SHEET 3: WEEKLY LOG
# ─────────────────────────────────────────────────────────────────────────────

def build_weekly_log(wb):
    ws = wb.create_sheet("Weekly Log")
    ws.sheet_properties.tabColor = "FB8C00"
    ws.sheet_view.showGridLines = False

    widths = [6, 12, 18, 28, 35, 35, 28, 9, 9, 9, 9, 9, 9, 11, 14, 30]
    for i, w in enumerate(widths, 1):
        cw(ws, i, w)

    rh(ws, 1, 42)
    merge_write(ws, 1, 1, 1, 16, "Weekly Log — Track What You Do Each Week",
                bold=True, size=18, fc="FFFFFF", bg="E65100", ha="center")
    rh(ws, 2, 8)

    # ── Column headers ─────────────────────────────────────────────────────
    rh(ws, 3, 28)
    headers_row1 = ["Wk", "Dates", "Phase", "Weekly Goal", "What I Completed",
                    "What I'm Stuck On", "Key Learning", "Mon", "Tue", "Wed", "Thu", "Fri", "Sat",
                    "Total Hrs", "Status", "Notes"]
    for col_i, h in enumerate(headers_row1, 1):
        write(ws, 3, col_i, h, bold=True, size=10, fc="FFFFFF", bg="37474F", ha="center", bc="546E7A")

    ws.freeze_panes = "A4"

    week_phases = [
        (1,"14 Jul – 20 Jul","Phase 1 — Foundation"),
        (2,"21 Jul – 27 Jul","Phase 1 — Foundation"),
        (3,"28 Jul – 3 Aug", "Phase 1 — Foundation"),
        (4,"4 Aug – 10 Aug", "Phase 2 — Backend"),
        (5,"11 Aug – 17 Aug","Phase 2 — Backend"),
        (6,"18 Aug – 24 Aug","Phase 2 — Backend"),
        (7,"25 Aug – 31 Aug","Phase 3 — Full Stack"),
        (8,"1 Sep – 7 Sep",  "Phase 3 — Full Stack"),
        (9,"8 Sep – 14 Sep", "Phase 3 — Full Stack"),
        (10,"15 Sep – 21 Sep","Phase 4 — GenAI"),
        (11,"22 Sep – 28 Sep","Phase 4 — GenAI"),
        (12,"29 Sep – 5 Oct", "Phase 4 — GenAI"),
        (13,"6 Oct – 12 Oct", "Phase 5 — Interview Project"),
        (14,"13 Oct – 19 Oct","Phase 5 — Interview Project"),
        (15,"20 Oct – 26 Oct","Phase 5 — Interview Project"),
        (16,"27 Oct – 2 Nov", "Phase 5 — Interview Project"),
    ]

    phase_color_map = {ph["name"]: ph for ph in PHASES}

    for row_offset, (wk, dates, phase_name) in enumerate(week_phases):
        r = 4 + row_offset
        rh(ws, r, 40)
        ph = phase_color_map[phase_name]
        hc = ph["color"]
        bg = ph["light"]

        write(ws, r,  1, f"W{wk}",     bold=True, size=11, fc="FFFFFF", bg=hc, ha="center", bc=hc)
        write(ws, r,  2, dates,         bold=False, size=9,  fc="333333", bg=bg, ha="center", bc="CCCCCC")
        write(ws, r,  3, phase_name,    bold=True,  size=9,  fc="FFFFFF", bg=hc, ha="left",   bc=hc)
        write(ws, r,  4, "",            bold=False, size=10, fc="333333", bg="FAFAFA", ha="left",   bc="CCCCCC")
        write(ws, r,  5, "",            bold=False, size=10, fc="1B5E20", bg="F1F8E9", ha="left",   bc="CCCCCC")
        write(ws, r,  6, "",            bold=False, size=10, fc="B71C1C", bg="FFF3E0", ha="left",   bc="CCCCCC")
        write(ws, r,  7, "",            bold=False, size=10, fc="333333", bg="E8EAF6", ha="left",   bc="CCCCCC")
        for col_i in range(8, 14):
            write(ws, r, col_i, None, bold=False, size=10, fc="333333", bg="FFFFFF", ha="center", bc="CCCCCC", num_fmt="0.0")
        # Total hours formula
        tot_cell = ws.cell(row=r, column=14)
        tot_cell.value = f"=SUM(H{r}:M{r})"
        tot_cell.font = font(bold=True, size=11, color="1976D2")
        tot_cell.alignment = align("center", "center")
        tot_cell.fill = fill("E3F2FD")
        tot_cell.border = border("CCCCCC")
        write(ws, r, 15, "Not Started", bold=False, size=10, fc="777777", bg="F5F5F5", ha="center", bc="CCCCCC")
        write(ws, r, 16, "",            bold=False, size=9,  fc="555555", bg="FFFDE7", ha="left",   bc="CCCCCC")

    add_dropdown(ws, '"Not Started,In Progress,Done,Skipped"', "O4:O19")

    # ── Total hours summary ─────────────────────────────────────────────────
    rh(ws, 21, 10)
    rh(ws, 22, 32)
    write(ws, 22, 1,  "TOTAL HOURS LOGGED", bold=True, size=12, fc="FFFFFF", bg="37474F", ha="left", bc="546E7A")
    tot_all = ws.cell(row=22, column=14)
    tot_all.value = "=SUM(N4:N19)"
    tot_all.font  = font(bold=True, size=14, color="76FF03")
    tot_all.fill  = fill("263238")
    tot_all.alignment = align("center", "center")
    tot_all.border = border("546E7A")

    return ws

# ─────────────────────────────────────────────────────────────────────────────
# SHEET 4: PROJECTS
# ─────────────────────────────────────────────────────────────────────────────

def build_projects(wb):
    ws = wb.create_sheet("Projects")
    ws.sheet_properties.tabColor = "8E24AA"
    ws.sheet_view.showGridLines = False

    cw(ws, 1,  3);  cw(ws, 2,  28); cw(ws, 3,  12); cw(ws, 4,  35)
    cw(ws, 5,  4);  cw(ws, 6,  28); cw(ws, 7,  12); cw(ws, 8,  35)

    rh(ws, 1, 42)
    merge_write(ws, 1, 1, 1, 8, "Project Tracker — TODO App  &  Document QnA Chatbot",
                bold=True, size=18, fc="FFFFFF", bg="4A148C", ha="center")

    # Project headers
    rh(ws, 2, 8)
    rh(ws, 3, 32)
    merge_write(ws, 3, 2, 3, 4, "PROJECT 1 — TODO App (Full Stack + AI)",
                bold=True, size=14, fc="FFFFFF", bg="1976D2")
    merge_write(ws, 3, 6, 3, 8, "PROJECT 2 — Document QnA Chatbot",
                bold=True, size=14, fc="FFFFFF", bg="3949AB")

    rh(ws, 4, 26)
    for col_i, (h, bg) in enumerate([("Feature / Milestone","1565C0"),("Status","1565C0"),("Notes","1565C0")], 2):
        write(ws, 4, col_i, h, bold=True, size=10, fc="FFFFFF", bg=bg, ha="center", bc="0D47A1")
    for col_i, (h, bg) in enumerate([("Feature / Milestone","283593"),("Status","283593"),("Notes","283593")], 6):
        write(ws, 4, col_i, h, bold=True, size=10, fc="FFFFFF", bg=bg, ha="center", bc="1A237E")

    # TODO App features
    todo_features = [
        ("Backend: Node.js + Express API",        ""),
        ("PostgreSQL DB with Prisma ORM",          ""),
        ("JWT Authentication (register + login)",  ""),
        ("Full CRUD: Tasks (protected routes)",    ""),
        ("Frontend: React + Vite",                 ""),
        ("Tailwind CSS styling",                   ""),
        ("Axios API integration",                  ""),
        ("Login / Register UI",                    ""),
        ("Protected routes (React Router v6)",     ""),
        ("Task list + CRUD UI",                    ""),
        ("AI: Ollama LLM integration",             ""),
        ("FAISS vector store",                     ""),
        ("LangChain RAG pipeline",                 ""),
        ("Natural language task input",            ""),
        ("AI smart suggestions feature",           ""),
    ]
    docqa_features = [
        ("Express.js backend project setup",       ""),
        ("File upload (Multer — PDF/TXT)",         ""),
        ("PDF/TXT parsing",                        ""),
        ("Text chunking (LangChain splitter)",     ""),
        ("Embeddings (nomic-embed-text)",          ""),
        ("FAISS vector store per document",        ""),
        ("RAG retrieval chain (LangChain JS)",     ""),
        ("Ollama LLM response generation",         ""),
        ("Streaming API response (SSE)",           ""),
        ("React chat UI (bubble thread)",          ""),
        ("File upload panel (drag & drop)",        ""),
        ("Streaming typewriter effect in UI",      ""),
        ("Loading + error states",                 ""),
        ("README with architecture diagram",       ""),
        ("Full local end-to-end test",             ""),
    ]

    for offset, (todo, docqa) in enumerate(zip(todo_features, docqa_features)):
        r = 5 + offset
        rh(ws, r, 30)
        bg1 = "E3F2FD" if offset % 2 == 0 else "FFFFFF"
        bg2 = "E8EAF6" if offset % 2 == 0 else "FFFFFF"
        write(ws, r, 2, todo[0],  bold=False, size=10, fc="222222", bg=bg1, ha="left",   bc="CCCCCC")
        write(ws, r, 3, "",       bold=False, size=10, fc="333333", bg=bg1, ha="center", bc="CCCCCC")
        write(ws, r, 4, "",       bold=False, size=9,  fc="555555", bg="FAFAFA", ha="left", bc="CCCCCC")
        write(ws, r, 6, docqa[0], bold=False, size=10, fc="222222", bg=bg2, ha="left",   bc="CCCCCC")
        write(ws, r, 7, "",       bold=False, size=10, fc="333333", bg=bg2, ha="center", bc="CCCCCC")
        write(ws, r, 8, "",       bold=False, size=9,  fc="555555", bg="FAFAFA", ha="left", bc="CCCCCC")

    # Dropdowns
    add_dropdown(ws, '"Not Started,In Progress,Done,Blocked"', f"C5:C{5+len(todo_features)-1}")
    add_dropdown(ws, '"Not Started,In Progress,Done,Blocked"', f"G5:G{5+len(docqa_features)-1}")

    return ws

# ─────────────────────────────────────────────────────────────────────────────
# SHEET 5: INTERVIEW PREP
# ─────────────────────────────────────────────────────────────────────────────

def build_interview_prep(wb):
    ws = wb.create_sheet("Interview Prep")
    ws.sheet_properties.tabColor = "3949AB"
    ws.sheet_view.showGridLines = False

    widths = [4, 22, 28, 14, 38, 38, 18, 35]
    for i, w in enumerate(widths, 1):
        cw(ws, i, w)

    rh(ws, 1, 42)
    merge_write(ws, 1, 1, 1, 8, "Interview Prep — Topics, Practice Questions & Confidence Tracker",
                bold=True, size=17, fc="FFFFFF", bg="1A237E", ha="center")

    rh(ws, 2, 8)
    rh(ws, 3, 30)
    headers = ["#", "Category", "Topic", "Confidence (1-5)", "Practice Question 1", "Practice Question 2", "Mock Score (1-10)", "Notes / Answers"]
    for col_i, h in enumerate(headers, 1):
        write(ws, 3, col_i, h, bold=True, size=10, fc="FFFFFF", bg="283593", ha="center", bc="1A237E")

    ws.freeze_panes = "A4"

    topics = [
        # Category, Topic, Q1, Q2
        ("System Design","Full Stack Architecture","Walk me through the architecture of your TODO app — frontend to DB.","Why did you choose PostgreSQL over MongoDB for this project?"),
        ("System Design","Database Design","Explain the schema — tables, relations, constraints you used.","What indexes did you add and why?"),
        ("Backend","Express.js Middleware","What is middleware? How does the middleware chain work?","How did you handle authentication in your Express app?"),
        ("Backend","JWT Authentication","Explain the full JWT auth flow — register, login, protected route.","What are the trade-offs between JWT and session-based auth?"),
        ("Backend","Prisma ORM","What is an ORM? Why did you use Prisma over raw SQL?","How does Prisma migration work?"),
        ("Frontend","React Hooks","Explain useState and useEffect with a real example from your project.","What is a custom hook and why did you create one?"),
        ("Frontend","State Management","How do you manage auth state across your React app?","When would you use Context API vs Redux?"),
        ("Frontend","API Integration","How does your React app communicate with the backend?","How did you handle token expiry on the frontend?"),
        ("GenAI","What is RAG?","Explain Retrieval-Augmented Generation in simple terms.","Why is RAG better than just prompting the LLM directly?"),
        ("GenAI","Embeddings","What are text embeddings? How do you generate them in your project?","What is cosine similarity and how is it used in your system?"),
        ("GenAI","FAISS","What is FAISS? Why did you choose it over other vector stores?","How does similaritySearch work in your Document QnA app?"),
        ("GenAI","LangChain JS","What role does LangChain play in your system?","Explain the retrieval chain you built step by step."),
        ("GenAI","Ollama","What is Ollama and why did you run the LLM locally?","How would you replace Ollama with a cloud LLM like OpenAI?"),
        ("Deployment","Vercel + Render","Explain your deployment architecture — what runs where?","How do frontend and backend communicate in production?"),
        ("Deployment","Cloud Database","Why did you switch from local PostgreSQL to Neon/Supabase?","How did you handle the migration to the cloud DB?"),
        ("Soft Skills","About the Project","Tell me about your Document QnA project — what does it do?","What was the hardest part of building it and how did you solve it?"),
        ("Soft Skills","Tech Decisions","Why did you choose Node.js over Python for your backend?","What would you do differently if you rebuilt this project?"),
        ("Soft Skills","Learning Journey","How did you learn all of this? What resources did you use?","Where do you see yourself in 2 years as a Full Stack + AI engineer?"),
        ("DSA","Arrays & Strings","Find the two numbers in an array that sum to a target.","Reverse words in a sentence without using split()."),
        ("DSA","Hash Maps","Find the first non-repeating character in a string.","Group anagrams from an array of strings."),
    ]

    for offset, (cat, topic, q1, q2) in enumerate(topics):
        r = 4 + offset
        rh(ws, r, 40)
        bg = "E8EAF6" if offset % 2 == 0 else "FFFFFF"
        write(ws, r, 1, str(offset+1), bold=True, size=10, fc="FFFFFF", bg="3949AB", ha="center", bc="283593")
        write(ws, r, 2, cat,   bold=True,  size=10, fc="FFFFFF", bg="3949AB", ha="left",   bc="283593")
        write(ws, r, 3, topic, bold=True,  size=10, fc="222222", bg=bg,       ha="left",   bc="CCCCCC")
        write(ws, r, 4, "",    bold=False, size=11, fc="333333", bg=bg,       ha="center", bc="CCCCCC")  # confidence
        write(ws, r, 5, q1,    bold=False, size=9,  fc="333333", bg="FAFAFA", ha="left",   bc="CCCCCC")
        write(ws, r, 6, q2,    bold=False, size=9,  fc="333333", bg="FAFAFA", ha="left",   bc="CCCCCC")
        write(ws, r, 7, "",    bold=False, size=11, fc="333333", bg=bg,       ha="center", bc="CCCCCC")  # mock score
        write(ws, r, 8, "",    bold=False, size=9,  fc="444444", bg="FFFDE7", ha="left",   bc="CCCCCC")

    add_dropdown(ws, '"1,2,3,4,5"',   f"D4:D{4+len(topics)-1}")
    add_dropdown(ws, '"1,2,3,4,5,6,7,8,9,10"', f"G4:G{4+len(topics)-1}")

    # Conditional formatting on confidence column
    ws.conditional_formatting.add(
        f"D4:D{4+len(topics)-1}",
        Rule(type="expression", formula=["D4>=4"], dxf=DifferentialStyle(fill=PatternFill(bgColor="C8E6C9", fill_type="solid"), font=Font(color="1B5E20")), stopIfTrue=True)
    )
    ws.conditional_formatting.add(
        f"D4:D{4+len(topics)-1}",
        Rule(type="expression", formula=["D4<=2"], dxf=DifferentialStyle(fill=PatternFill(bgColor="FFCDD2", fill_type="solid"), font=Font(color="B71C1C")), stopIfTrue=True)
    )

    return ws

# ─────────────────────────────────────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────────────────────────────────────

def main():
    wb = Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    print("Building Dashboard...")
    build_dashboard(wb)

    print("Building Task Tracker...")
    build_task_tracker(wb)

    print("Building Weekly Log...")
    build_weekly_log(wb)

    print("Building Projects...")
    build_projects(wb)

    print("Building Interview Prep...")
    build_interview_prep(wb)

    wb.save(OUTPUT)
    print(f"\n✅  Saved: {OUTPUT}")
    print(f"   Sheets: {[s.title for s in wb.worksheets]}")
    print(f"   Tasks:  {len(TASKS)} tasks across {len(PHASES)} phases")

if __name__ == "__main__":
    main()
